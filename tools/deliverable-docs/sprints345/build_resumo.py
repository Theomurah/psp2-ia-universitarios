#!/usr/bin/env python3
"""
Síntese do CORPO do artigo ENEGEP em ~7 páginas, preservando todas as
informações (todas as seções, métricas, citações e referências) — apenas
condensando a prosa, com uso de tópicos para densidade. Para análise.

Fonte: Entregas/Sprint 0/PSP2_Artigo_ENEGEP_Introducao.docx (15 págs).
Saída: Entregas/Sprint 0/PSP2_Artigo_Resumo_7paginas.docx
Rodar: python3 tools/deliverable-docs/sprints345/build_resumo.py
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
OUT = os.path.join(ROOT, 'Entregas', 'Sprint 0', 'PSP2_Artigo_Resumo_7paginas.docx')
FONT = 'Times New Roman'
BLUE = RGBColor(0x1F, 0x38, 0x64)

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.5); sec.left_margin = Cm(3.0); sec.bottom_margin = Cm(2.0); sec.right_margin = Cm(2.0)
normal = doc.styles['Normal']; normal.font.name = FONT; normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
pf = normal.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(4)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def H(text, size=13, before=8, after=3):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = BLUE; r.font.name = FONT
    return p


def P(text, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    # suporta **negrito** simples
    parts = text.split('**')
    for i, seg in enumerate(parts):
        if seg == '': continue
        r = p.add_run(seg); r.bold = (i % 2 == 1); r.font.name = FONT
    return p


def B(items, after=2):
    for it in items:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.space_after = Pt(after)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run('•  '); r.font.name = FONT
        parts = it.split('**')
        for i, seg in enumerate(parts):
            if seg == '': continue
            rr = p.add_run(seg); rr.bold = (i % 2 == 1); rr.font.name = FONT


# ===== Título =====
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('PSP2 — IA para Universitários'); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = BLUE; r.font.name = FONT
st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER; st.paragraph_format.space_after = Pt(6)
r = st.add_run('Síntese do artigo (corpo) para análise — preserva todas as informações; apenas condensada'); r.italic = True; r.font.size = Pt(10.5); r.font.name = FONT

# ===== 1. Introdução =====
H('1. Introdução')
H('1.1 Contextualização', size=11.5, before=4)
P('A IA generativa (ChatGPT, Claude, Gemini) transformou o cenário acadêmico; estudantes a usam para revisão, planos de estudo e preparação para avaliações. Porém, por padrão, essas plataformas operam de forma **genérica**, sem acesso ao contexto do semestre, disciplinas, materiais dos professores ou cronograma do aluno. Na UnB, um graduando gerencia tipicamente **6 a 9 disciplinas/semestre**, cada uma com material heterogêneo (PDF, DOCX, PPTX, imagens de quadro, Markdown) disperso em **SIGAA, e-mail, WhatsApp e pastas locais**, sem padrão de organização. Lacuna central: a IA está disponível, mas é ineficaz sem contexto personalizado.')
H('1.2 Problema', size=11.5, before=4)
P('Desconexão entre os materiais do aluno e a capacidade da IA de usá-los de forma contextualizada. Três obstáculos interdependentes:')
B(['Documentos **dispersos**, em múltiplos formatos e locais, sem padronização;',
   'Transformá-los em material organizado/sintético é **manual, demorado e sujeito a perdas**;',
   'Mesmo organizados, o aluno **não sabe configurar a IA** (falta engenharia de prompts e estruturação de contexto).'])
P('Resultado: uso **subótimo** (perguntas genéricas → respostas genéricas, desconectadas do que o professor cobra). Estima-se **3–5 h/semana** apenas organizando materiais.')
H('1.3 Justificativa', size=11.5, before=4)
B(['**Demanda observada:** levantamento informal com alunos de Eng. de Produção da UnB — ~**85%** usam IA generativa; **nenhum** configurou contexto personalizado. Barreira não é acesso, e sim a ausência de processo estruturado.',
   '**Viabilidade técnica:** Claude Projects, Custom GPTs e Gemini Gems já permitem conectar documentos e instruções — falta um **sistema intermediário** que prepare os documentos brutos.',
   '**Impacto:** reduzir em **50–60%** o tempo de organização e transformar uma LLM genérica em assistente calibrado ao contexto do aluno.'])
H('1.4 Objetivos', size=11.5, before=4)
P('**Geral:** desenvolver sistema web que recebe documentos brutos multi-formato, processa/sintetiza via LLM, organiza e exporta ao Google Drive com nomenclatura padronizada, e gera system prompt personalizado + biblioteca de prompts.')
P('**Específicos:** (a) API que recebe/processa **≥5 formatos** (PDF, DOCX, PPTX, imagens, Markdown); (b) pipeline LLM que classifica, sintetiza e comprime **sem perda de conteúdo técnico**; (c) integração com a API do Drive (pastas por disciplina); (d) gerador de system prompt + **≥8 prompts**; (e) interface web (upload, acompanhamento, visualização); (f) validar com **≥10 alunos** da UnB (satisfação, redução de tempo).')
P('**SMART:** **S** — protótipo funcional (recebe docs, processa via LLM, exporta organizado ao Drive, gera system prompt + biblioteca); **M** — classificação correta ≥80% em 50 docs; satisfação ≥70% com ≥10 alunos; redução ≥50% de tempo (manual vs sistema); **A** — equipe de 5, APIs em tier gratuito/freemium (Claude Haiku, Drive) e infra sem custo (Vercel, Supabase), em **5 sprints**; **R** — resolve desorganização, perda de tempo e falta de contexto; **T** — início 09/04/2026, entrega final (artigo + apresentação) até **25/06/2026**; avaliação empírica como etapa de continuidade.')

# ===== 2. Revisão Bibliográfica =====
H('2. Revisão Bibliográfica')
H('2.1 IA generativa no ensino superior', size=11.5, before=4)
P('Desde o ChatGPT (nov/2022), adoção de LLMs entre **80% e 90%** dos universitários (Sirnoorkar; Rebello, 2025; Hao et al., 2025). A literatura converge: uso **sem engenharia de prompts** gera resultados subótimos (Beale, 2025; Herklotz et al., 2025). Beale (2025) chama de **lacuna pedagógica** (eficácia depende de o usuário fornecer contexto). Herklotz et al. (2025): em 2.389 feedbacks de LLM em provas de estatística, **~7% continham erros** — daí a necessidade de validação automática e ancoragem curricular.')
H('2.2 Tutores inteligentes e RAG', size=11.5, before=4)
P('RAG (Lewis et al., 2020) conecta LLMs a bases externas, mitigando alucinações. Aplicações educacionais: **RAGMan** (Ma; Martins; Lopes, 2024) — 98% de precisão pedagógica em programação, 455 alunos; **CourseAssist** (Feng; Liu; Ghosal, 2024) — 6 disciplinas, 500+ alunos; **DeepTutor** (Zhao et al., 2026), **LeafTutor** (Bochard et al., 2025), **Stan** (Furst; Venkateshwaran, 2026) — tutores personalizados. Yu et al. (2024): MOOC → **MAIC**, agentes LLM elevam o engajamento. A maioria adota **arquitetura fechada** (Chen et al., 2025, RAG vs GraphRAG; Tufino, 2025, NotebookLM socrático).')
H('2.3 Lacunas e posicionamento', size=11.5, before=4)
B(['Sistemas mapeados são **plataformas fechadas**, sem artefato portátil reaproveitável;',
   'Produção **escassa** sobre pipelines que automatizam a etapa **anterior** à interação (transformar docs brutos em material pronto);',
   'Contexto **brasileiro sub-representado** na literatura internacional.'])
P('Posicionamento: gera saída **portátil** (system prompt + materiais no Drive) usável em qualquer LLM — proposta **complementar** que automatiza a preparação de contexto.')

# ===== 3. Metodologia =====
H('3. Metodologia')
B(['**3.1 Caracterização:** pesquisa **aplicada**; abordagem **mista** (quanti+quali; Creswell & Creswell, 2021); **exploratório-descritiva** (desenvolvimento) e **avaliativa** (testes). Método **Design Science Research** (Hevner et al., 2004; Peffers et al., 2007; Dresch; Lacerda; Antunes Jr., 2015). População: graduandos da UnB usuários de IA. Amostra: **10 por conveniência**, vários cursos (predomínio Eng. de Produção); exploratória, sem generalização estatística.',
   '**3.2 Contexto/participantes:** disciplina PSP2, 6º semestre de Eng. de Produção da UnB, 2026.1; voluntários recrutados em redes acadêmicas, com **TCLE** e anonimização; documentos isolados por **RLS** no banco.',
   '**3.3 Procedimentos:** projeto em **5 sprints** mapeadas às **6 atividades da DSRM** (Peffers et al., 2007): (i) identificação do problema; (ii) objetivos; (iii) projeto/desenvolvimento; (iv) demonstração; (v) avaliação; (vi) comunicação. (i)–(ii) = diagnóstico (Seção 1) + backlog; (iii)–(iv) = desenvolvimento e demonstração; (v) = validação técnica + avaliação com 10 alunos (**condicionada ao deploy, não concluída**); (vi) = artigo, apresentação e pacote de entrega.'])
H('3.4 Coleta de dados e instrumentos', size=11.5, before=4)
B(['**Frente 1 — validação técnica** (atividade 5.1): **50 documentos-teste** em 5 formatos; métricas: taxa de classificação correta, taxa de aprovação na validação em 4 camadas, **custo médio/doc (USD)** e **latência p50/p95**.',
   '**Frente 2 — avaliação com usuários** (atividade 5.2): sessões individuais de ~**60 min**, remotas; roteiro: cadastro → onboarding (curso, semestre, matérias, horários) → upload de 3 docs → revisão → uso do system prompt na IA de escolha. Instrumentos: (a) **SUS** (10 itens Likert-5, 0–100; Brooke, 1996; Bangor; Kortum; Miller, 2008); (b) **TAM** (Utilidade e Facilidade percebidas; Davis, 1989; Venkatesh; Davis, 2000); (c) **comparativo cronometrado** manual vs sistema; (d) **entrevista semiestruturada**.'])
H('3.5 Análise e métricas', size=11.5, before=4)
P('Quantitativa: estatística descritiva (média, mediana, desvio-padrão, p50/p95). **Critérios de aceitação:** (i) classificação **≥80%**; (ii) **SUS ≥70** (usabilidade boa; Bangor et al., 2008); (iii) satisfação **≥70%**; (iv) redução **≥50%** de tempo. Qualitativa: **análise de conteúdo** (Bardin, 2011), categorias emergentes. Limitações reconhecidas: amostra reduzida, conveniência e recorte à UnB → caráter exploratório.')

# ===== 4. Desenvolvimento do artefato =====
H('4. Desenvolvimento do artefato')
H('4.1 Arquitetura', size=11.5, before=4)
P('Web em **três camadas**: apresentação (SPA Vite + React 19 + TypeScript na **Vercel**); aplicação (**Supabase Edge Functions** em Deno, limite de **60 s** → pipeline **assíncrono**, status via **Realtime** do PostgreSQL); persistência (**PostgreSQL** com **RLS** em todas as tabelas de usuário). LLM via **OpenRouter** (intercambia Anthropic/OpenAI/Google por configuração, sem alterar código). Drive via **OAuth 2.0 escopo restrito** (drive.file — menor privilégio).')
H('4.2 Pipeline (5 estágios)', size=11.5, before=4)
B(['**Parsing:** 5 formatos (libs TS p/ texto/estruturados + modelo de visão p/ imagens) → texto bruto unificado;',
   '**Classificação:** LLM identifica matéria, tipo, data e título; saída via **Zod** com **confiança 0–1** (<0,4 → revisão manual; 0,4–0,7 → confirmação; >0,7 → automático);',
   '**Síntese:** prompt + contexto da classificação + perfil do aluno → Markdown com cabeçalho, **LaTeX**, tabelas e seção de erros comuns; compressão típica **20–50%**;',
   '**Compressão:** modo **compacto** (50–70%, revisão) e modo **cola** (15–25%, só fórmulas/tabelas/bullets) — mesmo template parametrizado;',
   '**Nomenclatura + exportação:** nome padronizado (MATÉRIA - Tipo - Identificador Título.ext) e export hierárquico ao Drive (semestre/matéria).'])
H('4.3 Validação em 4 camadas (custo crescente)', size=11.5, before=4)
B(['**Estrutural** (conformidade com schema Zod); **Quantitativa** (tamanho, nº de fórmulas, nº de seções); **Semântica** (keywords do original presentes na síntese → retenção); **LLM-as-judge** (2º modelo, distinto, nota **0–10** + justificativa).',
   'Decisões: **aprovação** / **retentativa** com prompt ajustado (até **2 iterações**) / **revisão pelo usuário**.'])
H('4.4 Artefato portátil, robustez e jornada', size=11.5, before=4)
P('**Saída portátil:** (1) estrutura hierárquica de pastas no Drive (organizada e nomeada); (2) **system prompt personalizado** (a partir das matérias, dos documentos e da biblioteca) inserível em **Custom GPTs, Claude Projects ou Gemini Gems** — calibra qualquer LLM ao contexto do aluno.')
P('**Robustez e segurança** (a partir de 4 rodadas de auditoria): acesso isolado por **RLS** + gatilhos **anti-escalada**; tokens de integração **fora do navegador**; conteúdo do aluno em **ambiente isolado** contra injeção de prompt; processamento assíncrono **resiliente** (corrida, reprocessamento, falhas parciais); **logs estruturados sem PII**; interface em **WCAG 2.1 AA**, responsiva e com tema escuro.')
P('**Jornada do aluno:** cadastro guiado (curso, semestre, disciplinas, horários — inclusive **importação do atestado do SIGAA**) sob **consentimento LGPD**; acompanhamento do processamento em tempo real; leitura das sínteses; gestão de histórico; **exportação e exclusão** de dados (direitos LGPD).')
P('**Biblioteca:** 8 prompts oficiais — resumo de aula, geração de questões de revisão, explicação simplificada, comparação entre conceitos, resolução de exercícios passo a passo, fichamento de capítulo, estudo dirigido e quiz pré-prova — copiáveis para uso externo.')

# ===== 5. Resultados e Discussão =====
H('5. Resultados e Discussão')
H('5.1 Artefato desenvolvido e demonstração', size=11.5, before=4)
P('Coerente com a DSR, o **principal resultado é o próprio artefato**: protótipo funcional que percorre ponta a ponta ingestão → classificação → síntese → compressão → exportação, **demonstrado em ambiente de desenvolvimento** nos 5 formatos, gerando a saída portátil. A **Tabela 1** sintetiza 9 componentes, todos **Implementados**: parsing multiformato; classificação por LLM; síntese; compressão (compacto/cola); validação em 4 camadas; geração de system prompt; biblioteca de 8 prompts; exportação para o Drive (OAuth drive.file); interface web. **A mensuração sistemática** (classificação, custo, latência sobre corpus controlado) foi **desenhada (Seção 3) mas não executada** — indicadores ainda não apurados; sua obtenção é etapa imediata de continuidade.')
H('5.2 Avaliação planejada e estado de execução', size=11.5, before=4)
P('A avaliação com usuários foi **integralmente planejada** (10 voluntários de cursos variados, sessões ~60 min, SUS + TAM + comparativo cronometrado + entrevista, critérios pré-definidos na Seção 3.5). Sua execução **dependia da disponibilização pública (deploy)**, não concluída na janela do ciclo. Logo, a coleta **não foi realizada** e **nenhum resultado empírico** (usabilidade, aceitação, redução de tempo) é reportado — optou-se **deliberadamente por não apresentar valores estimados/simulados**, preservando a fidedignidade. Roteiro de entrevista e protocolo de análise de conteúdo (Bardin, 2011) já preparados (categorias previstas: ganho de tempo, qualidade da síntese, portabilidade), prontos para aplicação imediata após a publicação.')
H('5.3 Discussão', size=11.5, before=4)
P('Sem dados próprios, a discussão se dá no plano do projeto e da aderência à literatura. RAGMan (Ma; Martins; Lopes, 2024), Stan (Furst; Venkateshwaran, 2026) e CourseAssist (Feng; Liu; Ghosal, 2024) reportam alta adequação pedagógica e boa usabilidade — **parâmetros de referência** futuros. A motivação dialoga com **Beale (2025)** (a barreira de adoção é preparar o contexto, não a tecnologia); o ganho mensurável (comparativo cronometrado) **permanece hipótese a testar**, em diálogo com Silva e Kampff (2025). A **diferenciação portátil** é a contribuição distintiva (a verificar). CGI.br (2025): **70%** dos estudantes do ensino médio usam IA, mas só **~1/3** recebeu orientação estruturada — reforça a relevância de estruturar o uso sem aprisionar o aluno a uma plataforma.')

# ===== 6. Considerações finais =====
H('6. Considerações finais')
P('O trabalho **desenvolveu e desenhou a avaliação** de um sistema que transforma materiais brutos em artefatos portáteis (Markdown organizado no Drive + system prompt) para uso contextualizado da IA. Conduzido sob **DSR**, em 5 sprints ↔ 6 atividades da DSRM. **Resultado principal:** protótipo funcional completo, demonstrado em desenvolvimento. As atividades de avaliação (técnica + com usuários) foram **integralmente desenhadas** (critérios e instrumentos definidos) mas **não executadas** (dependiam do deploy). **Nenhum resultado empírico é reportado.**')
P('**Contribuições:** teórica — modelo de **preparação automatizada de contexto** (etapa anterior à interação, pouco explorada na literatura, centrada em tutores fechados); prática — **artefato funcional** adaptável a outros contextos e cursos.')
P('**Limitações:** (1) **ausência de avaliação empírica** (deploy não concluído); (2) desenho com **10 participantes por conveniência**, em uma instituição → exploratório; (3) **dependência de serviços externos** (OpenRouter, Drive, Supabase).')
P('**Trabalhos futuros (ordem de prioridade):** concluir o **deploy** e executar a avaliação desenhada (validação técnica + SUS/TAM/cronometrado/entrevistas com os 10); ampliar para **>50 estudantes** em múltiplas instituições, com análise **inferencial**; **integração nativa com SIGAA/Moodle/Canvas** (elimina upload manual); **RAG dinâmico** (system prompt consulta os documentos em tempo de execução); biblioteca de prompts **colaborativa** com professores; **estudo longitudinal** (≥2 semestres).')

# ===== Referências =====
H('Referências (26)')
refs = [
 'BANGOR, A.; KORTUM, P. T.; MILLER, J. T. An empirical evaluation of the System Usability Scale. Int. J. of Human-Computer Interaction, v. 24, n. 6, p. 574-594, 2008.',
 'BARDIN, L. Análise de conteúdo. São Paulo: Edições 70, 2011.',
 'BEALE, R. Dialogic pedagogy for large language models. arXiv preprint, 2025.',
 'BOCHARD, M. et al. LeafTutor: an AI agent for programming assignment tutoring. arXiv preprint, 2025.',
 'BROOKE, J. SUS: a quick and dirty usability scale. In: Usability evaluation in industry. London: Taylor & Francis, 1996. p. 189-194.',
 'CGI.br. Pesquisa TIC Educação 2024. São Paulo: NIC.br, 2025.',
 'CHEN, S. et al. Comparing RAG and GraphRAG for math textbook question answering. arXiv preprint, 2025.',
 'CRESWELL, J. W.; CRESWELL, J. D. Projeto de pesquisa: métodos qualitativo, quantitativo e misto. 5. ed. Porto Alegre: Penso, 2021.',
 'DAVIS, F. D. Perceived usefulness, perceived ease of use, and user acceptance of IT. MIS Quarterly, v. 13, n. 3, p. 319-340, 1989.',
 'DRESCH, A.; LACERDA, D. P.; ANTUNES JR., J. A. V. Design science research. Porto Alegre: Bookman, 2015.',
 'FENG, X.; LIU, R.; GHOSAL, T. CourseAssist: pedagogically appropriate AI tutor for CS education. arXiv preprint, 2024.',
 'FURST, A.; VENKATESHWARAN, S. Stan: an LLM-based thermodynamics course assistant. arXiv preprint, 2026.',
 'GIL, A. C. Métodos e técnicas de pesquisa social. 7. ed. São Paulo: Atlas, 2019.',
 'HAO, Z. et al. Student-AI interaction in an LLM-empowered learning environment. arXiv preprint, 2025.',
 'HERKLOTZ, M. et al. Can we trust LLMs as a tutor? Evaluating LLM-generated feedback in statistics exams. arXiv preprint, 2025.',
 'HEVNER, A. R. et al. Design science in information systems research. MIS Quarterly, v. 28, n. 1, p. 75-105, 2004.',
 'LACERDA, D. P. et al. Design science research: método de pesquisa para a engenharia de produção. Gestão & Produção, v. 20, n. 4, p. 741-761, 2013.',
 'LEWIS, P. et al. Retrieval-augmented generation for knowledge-intensive NLP tasks. NeurIPS, v. 33, p. 9459-9474, 2020.',
 'MA, J.; MARTINS, P.; LOPES, R. RAGMan: integrating AI tutors in programming course. arXiv preprint, 2024.',
 'PEFFERS, K. et al. A design science research methodology for IS research. J. of MIS, v. 24, n. 3, p. 45-77, 2007.',
 'SILVA, D. S. da; KAMPFF, A. J. C. Licenciatura em tempos de inteligência artificial. Rev. Docência do Ensino Superior, v. 15, p. 1-19, 2025.',
 'SIRNOORKAR, A.; REBELLO, N. S. Feedback that clicks: physics students valued features in AI feedback. arXiv preprint, 2025.',
 'TUFINO, E. NotebookLM as a Socratic physics tutor. arXiv preprint, 2025.',
 'VENKATESH, V.; DAVIS, F. D. A theoretical extension of the technology acceptance model. Management Science, v. 46, n. 2, p. 186-204, 2000.',
 'YU, J. et al. From MOOC to MAIC: reshaping online teaching through LLM-driven agents. arXiv preprint, 2024.',
 'ZHAO, B. et al. DeepTutor: towards agentic personalized tutoring. arXiv preprint, 2026.',
]
for rf in refs:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(rf); r.font.size = Pt(9.5); r.font.name = FONT

doc.save(OUT)
print('OK ->', os.path.relpath(OUT, ROOT))
