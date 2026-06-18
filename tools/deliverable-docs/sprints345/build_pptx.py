#!/usr/bin/env python3
"""
Deck de fechamento (Sprints 1-5) — estilo UnB do deck de referência, elevado:
fontes maiores (~+25%), cards editoriais com trilho de acento (sem caixa pesada),
acento amarelo UnB no header, número de slide, mais respiro. Rev. 3.

Rodar: python3 tools/deliverable-docs/sprints345/build_pptx.py
"""
import os, io, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
REF = os.path.join(ROOT, 'Entregas', 'PSP2_Apresentacao_Estado_do_Projeto.pptx')
OUT = os.path.join(ROOT, 'Entregas', 'PSP2_Apresentacao_Estado_do_Projeto_2026-06-18.pptx')

BLUE   = RGBColor(0x00, 0x33, 0x66)
BLUE_D = RGBColor(0x0A, 0x2A, 0x4A)
GREEN  = RGBColor(0x00, 0x85, 0x42)
GREEN_D= RGBColor(0x00, 0x59, 0x23)
YELLOW = RGBColor(0xFF, 0xB8, 0x1C)
DARK   = RGBColor(0x1F, 0x29, 0x37)
GRAY   = RGBColor(0x55, 0x60, 0x6E)
LGRAY  = RGBColor(0x8A, 0x93, 0x9F)
LBLUE  = RGBColor(0xBF, 0xD0, 0xE6)
BORDER = RGBColor(0xE6, 0xE9, 0xEE)
CARD   = RGBColor(0xF7, 0xF9, 0xFB)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
WARN   = RGBColor(0xB8, 0x5C, 0x00)
WARNBG = RGBColor(0xFB, 0xE9, 0xD3)
GREENBG= RGBColor(0xDD, 0xF0, 0xE3)
FONT = 'Calibri'
SW, SH = 13.333, 7.5
ML = 0.55
CW = SW - 2 * ML

_zip = zipfile.ZipFile(REF)
LOGO_WHITE = _zip.read('ppt/media/image-1-1.png')
LOGO_CONTENT = _zip.read('ppt/media/image-2-1.png')

prs = Presentation()
prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]
_pageno = [0]


def slide(bg=None):
    s = prs.slides.add_slide(BLANK); _pageno[0] += 1
    if bg is not None:
        cSld = s._element.find(qn('p:cSld'))
        bgEl = etree.SubElement(cSld, qn('p:bg')); bgPr = etree.SubElement(bgEl, qn('p:bgPr'))
        sf = etree.SubElement(bgPr, qn('a:solidFill')); clr = etree.SubElement(sf, qn('a:srgbClr'))
        clr.set('val', '%02X%02X%02X' % (bg[0], bg[1], bg[2])); etree.SubElement(bgPr, qn('a:effectLst'))
        cSld.insert(0, bgEl)
    return s


def _font(run, size, color, bold=False, italic=False):
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.name = FONT; run.font.color.rgb = color


def _fill_tf(tf, lines, align, anchor, sp_after, sp_line):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('left', 'right', 'top', 'bottom'): setattr(tf, 'margin_' + m, 0)
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.space_after = Pt(sp_after); p.space_before = Pt(0)
        if sp_line: p.line_spacing = sp_line
        for (t, sz, col, bd, *it) in (ln if isinstance(ln, list) else [ln]):
            r = p.add_run(); r.text = t; _font(r, sz, col, bd, it[0] if it else False)


def txt(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=3, sp_line=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tb.shadow.inherit = False
    _fill_tf(tb.text_frame, lines, align, anchor, sp_after, sp_line); return tb


def shape(s, kind, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None):
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)); sp.shadow.inherit = False
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    if radius is not None:
        try: sp.adjustments[0] = radius
        except Exception: pass
    return sp


def shtext(sp, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=1):
    _fill_tf(sp.text_frame, lines, align, anchor, sp_after, None)


def card(s, x, y, w, h, accent, fill=CARD, rail='left'):
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, line=BORDER, line_w=0.75)
    if rail == 'left':
        shape(s, MSO_SHAPE.RECTANGLE, x, y, 0.075, h, fill=accent, line=accent)
    elif rail == 'top':
        shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.075, fill=accent, line=accent)


def callout(s, x, y, w, h, fill, line):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, line_w=1.0, radius=0.06)


def header(s, eyebrow, kicker, title, title_size=37, title_h=1.05):
    w = max(1.6, 0.135 * len(eyebrow) + 0.5)
    pill = shape(s, MSO_SHAPE.RECTANGLE, ML, 0.32, w, 0.44, fill=BLUE, line=BLUE)
    shtext(pill, [[(eyebrow.upper(), 15.5, WHITE, True)]])
    s.shapes.add_picture(io.BytesIO(LOGO_CONTENT), Inches(12.5), Inches(0.18), Inches(0.66), Inches(0.66))
    # divisor: linha cinza + acento verde + tick amarelo (UnB)
    shape(s, MSO_SHAPE.RECTANGLE, ML, 1.02, CW, 0.018, fill=BORDER, line=None)
    shape(s, MSO_SHAPE.RECTANGLE, ML, 1.0, 1.5, 0.05, fill=GREEN, line=None)
    shape(s, MSO_SHAPE.RECTANGLE, ML, 1.0, 0.42, 0.05, fill=YELLOW, line=None)
    txt(s, ML, 1.16, CW, 0.34, [[(kicker.upper(), 16.5, GREEN, True)]])
    txt(s, ML, 1.52, CW, title_h, [[(title, title_size, BLUE, True)]], sp_line=0.95)
    pageno(s)


def pageno(s, dark=False):
    txt(s, SW - 1.1, SH - 0.46, 0.7, 0.3, [[('%02d' % _pageno[0], 11, (LBLUE if dark else LGRAY), False)]], align=PP_ALIGN.RIGHT)


def bullet(s, x, y, w, text_, size=18, color=DARK, dot=GREEN, sp_line=1.0):
    shape(s, MSO_SHAPE.OVAL, x, y + 0.09, 0.15, 0.15, fill=dot, line=dot)
    txt(s, x + 0.34, y, w - 0.34, 0.6, [[(text_, size, color, False)]], sp_line=sp_line)


def numbered(s, x, y, n, d=0.58):
    shtext(shape(s, MSO_SHAPE.OVAL, x, y, d, d, fill=BLUE, line=BLUE), [[(str(n), 19, WHITE, True)]])


def mantra(s, text_, y=6.95):
    shape(s, MSO_SHAPE.RECTANGLE, ML, y - 0.02, 0.6, 0.05, fill=GREEN, line=None)
    txt(s, ML, y + 0.09, CW - 1.2, 0.4, [[(text_, 17, BLUE, True)]])


def status(s, x, y, w, label, fg, bg):
    shtext(shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.34, fill=bg, line=bg), [[(label.upper(), 11.5, fg, True)]])


# ============================ S1 — CAPA ============================
s = slide(bg=BLUE)
s.shapes.add_picture(io.BytesIO(LOGO_WHITE), Inches(5.92), Inches(0.85), Inches(1.5), Inches(1.33))
txt(s, ML, 2.55, CW, 1.0, [[('PSP2 — IA para Universitários', 52, WHITE, True)]], align=PP_ALIGN.CENTER)
txt(s, ML, 3.62, CW, 0.55, [[('Estado do Projeto — Fechamento do ciclo (Sprints 1–5)', 25, LBLUE, False)]], align=PP_ALIGN.CENTER)
shape(s, MSO_SHAPE.RECTANGLE, 6.07, 4.4, 1.2, 0.05, fill=GREEN, line=None)
shape(s, MSO_SHAPE.RECTANGLE, 6.07, 4.4, 0.35, 0.05, fill=YELLOW, line=None)
txt(s, ML, 4.62, CW, 0.35, [[('EQUIPE', 17, GREEN, True)]], align=PP_ALIGN.CENTER)
txt(s, ML, 4.98, CW, 0.4, [[('Theo Murahovschi   ·   Pedro Henrique   ·   Isaac   ·   Guilherme   ·   Luis Felipe', 21, WHITE, False)]], align=PP_ALIGN.CENTER)
txt(s, ML, 6.5, CW, 0.35, [[('Universidade de Brasília   ·   Engenharia de Produção   ·   2026.1   ·   18/06/2026', 15, LBLUE, False)]], align=PP_ALIGN.CENTER)

# ============================ S2 — A IDEIA ============================
s = slide(); header(s, 'A Ideia', 'Visão geral · o que queremos resolver', 'Um hub que organiza o estudo do aluno UnB')
txt(s, ML, 2.7, 6.85, 1.4, [[('O aluno recebe material solto — PDFs, slides, fotos, DOCX — em cinco sistemas diferentes. ', 19, DARK, False), ('Resultado: perde tempo organizando em vez de estudar.', 19, DARK, True)]], sp_line=1.08)
txt(s, ML, 4.25, 6.85, 0.4, [[('O que ele recebe de volta:', 19, BLUE, True)]])
for i, b in enumerate(['Um resumo confiável de cada documento', 'Tudo nomeado e organizado por matéria, automaticamente', 'Uma biblioteca de prompts prontos pra qualquer IA', 'O próprio Drive do aluno como espelho do estudo']):
    bullet(s, ML + 0.05, 4.78 + i * 0.5, 6.8, b)
mantra(s, 'Menos tempo organizando. Mais tempo aprendendo.')
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.85, 2.7, 4.93, 4.0, fill=CARD, line=BORDER, line_w=0.75, radius=0.04)
txt(s, 8.08, 2.9, 4.5, 0.35, [[('Como funciona, em alto nível', 16, BLUE, True)]])
for i, (t, d) in enumerate([('O aluno envia', 'PDFs, slides, fotos, DOCX'), ('A IA entende', 'Identifica matéria e conteúdo'), ('A IA condensa', 'Sintetiza e gera a revisão'), ('O aluno usa', 'Drive + prompts personalizados')]):
    y = 3.38 + i * 0.82; numbered(s, 8.12, y, i + 1)
    txt(s, 8.85, y - 0.02, 3.8, 0.7, [[(t, 16, BLUE, True)], [(d, 13.5, GRAY, False)]], sp_after=1)

# ============================ S3 — REPLANEJAMENTO ============================
s = slide(); header(s, 'Replanejamento', 'Do plano inicial ao plano real', 'O plano agora reflete o que realmente foi feito')
txt(s, ML, 2.72, CW, 0.45, [[('A backlog inicial (4 sprints) foi reformulada retroativamente em 5 sprints — deadline 25/06 mantido.', 18, GRAY, False)]])
cwm = CW / 2 - 0.18
card(s, ML, 3.35, cwm, 2.55, RGBColor(0x9C, 0xA3, 0xAF), fill=CARD, rail='top')
status(s, ML + 0.28, 3.6, 1.9, 'Antes · 4 sprints', fg=GRAY, bg=RGBColor(0xEA, 0xEC, 0xEF))
txt(s, ML + 0.28, 4.18, cwm - 0.55, 1.6, [[('S3 · Integração + Testes com usuários', 17, DARK, False)], [('S4 · Feedback + Artigo + Apresentação', 17, DARK, False)]], sp_after=12)
dx = ML + CW / 2 + 0.18
card(s, dx, 3.35, cwm, 2.55, GREEN, fill=GREENBG, rail='top')
status(s, dx + 0.28, 3.6, 2.1, 'Depois · 5 sprints', fg=WHITE, bg=GREEN)
txt(s, dx + 0.28, 4.18, cwm - 0.55, 1.6, [[('S3 · Hardening completo do sistema', 17, GREEN_D, True)], [('S4 · Elaboração e finalização do artigo', 17, GREEN_D, True)], [('S5 · Jornada do cliente', 17, GREEN_D, True)]], sp_after=8)
txt(s, ML, 6.1, CW, 0.5, [[('O produto cresceu na prática (segurança, LGPD, jornada). O plano passou a capturar isso, em vez de tratar como "extra".', 16, GRAY, False)]])
mantra(s, 'Agora plano e produto contam a mesma história.')

# ============================ S4 — BACKLOG & ENTREGAS ============================
s = slide(); header(s, 'Backlog', 'Escopo completo em 5 sprints', 'Backlog do produto — 19 histórias, 5 sprints')
cwn = (CW - 3 * 0.22) / 4
for i, (n, l) in enumerate([('19', 'histórias'), ('69', 'tarefas'), ('5', 'sprints'), ('5', 'integrantes')]):
    x = ML + i * (cwn + 0.22); card(s, x, 2.62, cwn, 0.98, GREEN, rail='left')
    txt(s, x + 0.28, 2.62, cwn - 0.4, 0.98, [[(n, 42, GREEN, True), ('  ' + l, 15, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE)
sprints = [('SPRINT 1', 'Fundação', 'H1–H4 · 21 tarefas', 'CONCLUÍDA', GREEN),
           ('SPRINT 2', 'Inteligência', 'H5–H7 · 13 tarefas', 'CONCLUÍDA', GREEN),
           ('SPRINT 3', 'Hardening', 'H13–H18 · 18 tarefas', 'CONCLUÍDA', GREEN),
           ('SPRINT 4', 'Artigo', 'H19–H20 · 6 tarefas', 'EM FINALIZAÇÃO', WARN),
           ('SPRINT 5', 'Jornada do cliente', 'H21–H24 · 11 tarefas', 'CONCLUÍDA', GREEN)]
cw = (CW - 4 * 0.2) / 5
for i, (sp, nm, tk, st, c) in enumerate(sprints):
    x = ML + i * (cw + 0.2)
    shtext(shape(s, MSO_SHAPE.RECTANGLE, x, 3.85, cw, 0.5, fill=BLUE, line=BLUE), [[(sp, 14, WHITE, True)]])
    shape(s, MSO_SHAPE.RECTANGLE, x, 4.35, cw, 1.45, fill=WHITE, line=BORDER, line_w=0.75)
    txt(s, x + 0.15, 4.48, cw - 0.3, 1.1, [[(nm, 15, BLUE, True)], [(tk, 13, GRAY, False)]], sp_after=4)
    status(s, x + 0.15, 5.4, cw - 0.3, st, fg=c, bg=(GREENBG if c == GREEN else WARNBG))
callout(s, ML, 6.0, CW, 0.95, CARD, BORDER)
txt(s, ML + 0.3, 6.12, CW - 0.6, 0.75, [[('Entregas também atualizadas', 15, BLUE, True)], [('docs de tarefa (T35–T69) · relatórios por sprint (3 e 4 revisados, 5 novo) · relatório geral · artigo ENEGEP', 14, DARK, False)]], sp_after=2)

# ============================ S5 — FLUXO ============================
s = slide(); header(s, 'Fluxo', 'Do começo ao fim', 'As 5 sprints, lado a lado')
flow = [('SPRINT 1', 'Fundação', ['Arquitetura & stack', 'API de upload', 'Interface drag-and-drop'], 'CONCLUÍDA', GREEN),
        ('SPRINT 2', 'Inteligência', ['Pipeline LLM', 'Exportação Drive', 'System prompt + biblioteca'], 'CONCLUÍDA', GREEN),
        ('SPRINT 3', 'Hardening', ['Segurança & RLS', 'Pipeline resiliente', 'Observabilidade, a11y, CI/CD'], 'CONCLUÍDA', GREEN),
        ('SPRINT 4', 'Artigo', ['Corpo ABNT (6 seções)', 'Método DSR', 'Finalização & submissão'], 'EM FINALIZAÇÃO', WARN),
        ('SPRINT 5', 'Jornada', ['Onboarding & LGPD', 'Perfil + SIGAA + Drive', 'Jornada do cliente'], 'CONCLUÍDA', GREEN)]
cw = (CW - 4 * 0.2) / 5
for i, (sp, nm, items, st, c) in enumerate(flow):
    x = ML + i * (cw + 0.2)
    shtext(shape(s, MSO_SHAPE.RECTANGLE, x, 2.72, cw, 0.78, fill=BLUE, line=BLUE), [[(sp, 14, WHITE, True)], [(nm, 12.5, LBLUE, False)]], sp_after=0)
    shape(s, MSO_SHAPE.RECTANGLE, x, 3.5, cw, 2.5, fill=WHITE, line=BORDER, line_w=0.75)
    for j, it in enumerate(items):
        shape(s, MSO_SHAPE.OVAL, x + 0.16, 3.72 + j * 0.62, 0.12, 0.12, fill=GREEN, line=GREEN)
        txt(s, x + 0.36, 3.64 + j * 0.62, cw - 0.46, 0.6, [[(it, 12.5, DARK, False)]], sp_line=0.95)
    status(s, x + 0.16, 5.56, cw - 0.32, st, fg=c, bg=(GREENBG if c == GREEN else WARNBG))
mantra(s, 'Cinco etapas, uma linha contínua.', y=6.45)

# ============================ S6 — SPRINT 3 ============================
s = slide(); header(s, 'Fluxo real & robustez', 'Sprint 3 · Hardening', 'O sistema passou a funcionar de verdade')
cards6 = [('Login & autenticação', 'Entrada segura, sessão e proteção de rotas'),
          ('Conexão Google Drive', 'OAuth, refresh de token, upload idempotente'),
          ('Estruturação de pastas', 'Hierarquia por semestre e matéria'),
          ('Fluxo ponta-a-ponta', 'Do upload à exportação, sem passo manual'),
          ('Banco de dados sólido', 'RLS, anti-escalada, tokens fora do navegador')]
cw = (CW - 4 * 0.2) / 5
for i, (t, d) in enumerate(cards6):
    x = ML + i * (cw + 0.2); card(s, x, 2.62, cw, 2.0, GREEN, rail='top')
    txt(s, x + 0.16, 2.82, cw - 0.32, 1.7, [[(t, 14.5, BLUE, True)], [(d, 12.5, GRAY, False)]], sp_after=6, sp_line=0.98)
callout(s, ML, 4.85, CW, 1.42, GREENBG, GREEN)
txt(s, ML + 0.32, 4.98, CW - 0.64, 1.2, [
    [('Hardening guiado por 4 rodadas de auditoria — ~200 achados, 1 crítico (escalada de admin) resolvido.', 15.5, DARK, True)],
    [('E mais: correção de bugs · correção de falhas de segurança · otimizações de código · inovações (calendário, cronograma, biblioteca de prompts).', 14.5, GREEN_D, False)],
], sp_after=4, sp_line=1.0)
mantra(s, 'De peças soltas a um fluxo sólido.')

# ============================ S7 — SPRINT 4 (ARTIGO) ============================
s = slide(); header(s, 'Pesquisa', 'Sprint 4 · Artigo', 'Artigo ENEGEP 2026 — pronto e honesto')
b7 = [('Artigo científico completo sob o método Design Science Research, em 6 seções no padrão ABNT NBR 14724 (~14 páginas).', GREEN),
      ('O artefato — protótipo funcional de ponta a ponta — é o resultado principal do ciclo, sintetizado na Tabela 1 de componentes.', GREEN),
      ('Integridade metodológica: a avaliação com usuários (SUS, TAM) é reportada como planejada, não executada — sem dados fabricados.', WARN)]
for i, (t, c) in enumerate(b7):
    y = 2.78 + i * 1.2; card(s, ML, y, CW, 1.0, c, rail='left')
    txt(s, ML + 0.42, y, CW - 0.75, 1.0, [[(t, 17, DARK, False)]], anchor=MSO_ANCHOR.MIDDLE, sp_line=1.02)
mantra(s, 'Honestidade vale mais que número inventado.')

# ============================ S8 — SPRINT 5 (JORNADA) ============================
s = slide(); header(s, 'Experiência', 'Sprint 5 · Jornada do cliente', 'A jornada completa do aluno, ponta a ponta')
j = [('Aquisição & Onboarding', 'Login + magic link, identidade UnB, consentimento LGPD, onboarding em 4 passos'),
     ('Perfil acadêmico', 'Grade visual de horários + importação do SIGAA, conexão do Drive'),
     ('Uso diário', 'Dashboard em tempo real, gestão de documentos, transparência total'),
     ('Voz do cliente', 'Biblioteca de prompts + system prompt portátil, feedback, direitos LGPD')]
cw = (CW - 0.32) / 2
for i, (t, d) in enumerate(j):
    x = ML + (i % 2) * (cw + 0.32); y = 2.62 + (i // 2) * 1.62
    card(s, x, y, cw, 1.42, BLUE, rail='left')
    txt(s, x + 0.36, y + 0.2, cw - 0.6, 1.1, [[(t, 17, BLUE, True)], [(d, 14, GRAY, False)]], sp_after=5, sp_line=0.98)
callout(s, ML, 5.95, CW, 0.92, GREENBG, GREEN)
txt(s, ML + 0.32, 5.95, CW - 0.64, 0.92, [[('Mapeamos toda a jornada do cliente — da descoberta ao hábito de uso — para guiar as próximas decisões de produto.', 16, GREEN_D, True)]], anchor=MSO_ANCHOR.MIDDLE)
mantra(s, 'Da primeira tela ao hábito de estudo.', y=7.0)

# ============================ S9 — MATURIDADE & EXTRAS ============================
s = slide(); header(s, 'Além do escopo', 'O produto entregue é maior que o planejado', 'O que cresceu durante a execução')
ex = [('Experiência do aluno', 'Identidade UnB, onboarding, busca/filtros, arquivar/excluir, tema escuro'),
      ('Acompanhamento', 'Página de atividade, métricas pessoais, painel administrativo'),
      ('Segurança & privacidade', 'Hardening de acesso, sandbox anti-injeção, LGPD completa'),
      ('Inteligência do pipeline', 'Validação em 4 camadas (LLM-as-judge), import SIGAA, modos compacta/cola'),
      ('Confiabilidade', 'Logs sem PII, atualização em tempo real, CI/CD com gate, auditoria multi-frente'),
      ('Flexibilidade', 'Troca de modelo sem deploy, multi-provider de IA, biblioteca de prompts')]
cw = (CW - 2 * 0.25) / 3
for i, (t, d) in enumerate(ex):
    x = ML + (i % 3) * (cw + 0.25); y = 2.62 + (i // 3) * 1.85
    card(s, x, y, cw, 1.62, GREEN, rail='top')
    txt(s, x + 0.2, y + 0.2, cw - 0.4, 1.3, [[(t, 15.5, BLUE, True)], [(d, 13, GRAY, False)]], sp_after=5, sp_line=0.98)
mantra(s, 'O produto cresceu — e o plano acompanhou.')

# ============================ S10 — VISÃO DE FUTURO ============================
s = slide(bg=BLUE)
shtext(shape(s, MSO_SHAPE.RECTANGLE, ML, 0.5, 2.45, 0.46, fill=GREEN, line=GREEN), [[('ALÉM DA DISCIPLINA', 15.5, WHITE, True)]])
s.shapes.add_picture(io.BytesIO(LOGO_WHITE), Inches(12.3), Inches(0.42), Inches(0.66), Inches(0.58))
txt(s, ML, 1.2, CW, 0.85, [[('Sistema inteligente unificado do estudante', 38, WHITE, True)]])
txt(s, ML, 2.12, CW, 0.4, [[('O hub único onde o aluno UnB centraliza tudo que importa pro semestre.', 18, LBLUE, False)]])
cw = (CW - 0.4) / 2
for k, (head, items) in enumerate([
    ('TUDO EM UM LUGAR SÓ', ['Plataformas acadêmicas: universidade, fórum, Drive, e-mail e calendário', 'Caixa de entrada única: avisos, prazos e mensagens consolidados', 'Painel do semestre: matérias, notas, grade e progresso']),
    ('IA QUE CONHECE O ALUNO', ['Plano de estudo personalizado, sob medida pra carga do semestre', 'Reforço retroativo: identifica lacunas de semestres passados', 'Tutor pessoal por matéria, no nível do aluno'])]):
    x = ML + k * (cw + 0.4)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.8, cw, 3.0, fill=BLUE_D, line=RGBColor(0x2A, 0x52, 0x7E), line_w=0.75, radius=0.04)
    shape(s, MSO_SHAPE.RECTANGLE, x, 2.8, cw, 0.06, fill=YELLOW if k else GREEN, line=None)
    txt(s, x + 0.3, 3.0, cw - 0.6, 0.35, [[(head, 15, YELLOW if k else GREEN, True)]])
    for j, it in enumerate(items):
        shape(s, MSO_SHAPE.OVAL, x + 0.32, 3.62 + j * 0.66 + 0.04, 0.13, 0.13, fill=GREEN, line=GREEN)
        txt(s, x + 0.58, 3.56 + j * 0.66, cw - 0.9, 0.65, [[(it, 14, WHITE, False)]], sp_line=0.95)
txt(s, ML, 6.05, CW, 0.4, [[('Também à frente: ', 14, LBLUE, True), ('planejamento e produtividade · acessibilidade · privacidade avançada · monetização sustentável.', 14, LBLUE, False)]])
shape(s, MSO_SHAPE.RECTANGLE, ML, 6.6, 0.6, 0.05, fill=GREEN, line=None)
txt(s, ML, 6.7, CW, 0.4, [[('Menos abas abertas. Mais foco. Mais tempo pro aluno.', 17, WHITE, True)]])
pageno(s, dark=True)

# ============================ S11 — MACRO & MICRO ============================
s = slide(); header(s, 'Para onde vai', 'Da disciplina à pesquisa', 'Como seguimos — e por que vira PIBIC')
cw = (CW - 0.32) / 2
shtext(shape(s, MSO_SHAPE.RECTANGLE, ML, 2.62, cw, 0.5, fill=BLUE, line=BLUE), [[('MACRO · direção', 15.5, WHITE, True)]])
shape(s, MSO_SHAPE.RECTANGLE, ML, 3.12, cw, 3.45, fill=CARD, line=BORDER, line_w=0.75)
for j, (t, hl) in enumerate([('Testar a direção: features novas, testes, feedbacks e análise de dados', False), ('Pesquisa de mercado em paralelo', False), ('Transformar o processo em PIBIC — continuidade científica (ProIC/UnB)', True)]):
    shape(s, MSO_SHAPE.OVAL, ML + 0.25, 3.42 + j * 1.0 + 0.05, 0.15, 0.15, fill=GREEN, line=GREEN)
    txt(s, ML + 0.52, 3.36 + j * 1.0, cw - 0.8, 0.95, [[(t, 15.5, GREEN_D if hl else DARK, hl)]], sp_line=1.0)
mx = ML + cw + 0.32
shtext(shape(s, MSO_SHAPE.RECTANGLE, mx, 2.62, cw, 0.5, fill=GREEN, line=GREEN), [[('MICRO · próximas ações', 15.5, WHITE, True)]])
shape(s, MSO_SHAPE.RECTANGLE, mx, 3.12, cw, 3.45, fill=CARD, line=BORDER, line_w=0.75)
micro = ['Rodar a validação técnica (50 docs) → números reais pro artigo',
         'Entrar no PIBIC: edital, orientador(a), projeto + plano de trabalho',
         'Deploy em ambiente público',
         'Seleção e acompanhamento de usuários-teste',
         'Hardening contínuo: segurança, logs e banco',
         'Novas features para sentir a direção — ex.: flashcards']
for j, it in enumerate(micro):
    shape(s, MSO_SHAPE.OVAL, mx + 0.25, 3.4 + j * 0.55 + 0.03, 0.13, 0.13, fill=GREEN, line=GREEN)
    txt(s, mx + 0.5, 3.34 + j * 0.55, cw - 0.75, 0.55, [[(it, 13.5, DARK, False)]], sp_line=0.95)
mantra(s, 'A disciplina acaba; a pesquisa começa.')

prs.save(OUT)
print('OK ->', os.path.relpath(OUT, ROOT), '| slides:', len(prs.slides._sldIdLst))
